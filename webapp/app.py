import os
import matplotlib
matplotlib.use("Agg")  # MUY IMPORTANTE en Render
import matplotlib.pyplot as plt
from flask import Flask, render_template, request, redirect, url_for
from webapp.db import (
    crear_tablas,
    obtener_temporadas,
    insertar_temporada,
    obtener_temporada_por_id,
    obtener_equipos_por_temporada,
    insertar_equipo,
    obtener_equipo_por_id,
    obtener_jugadores_por_equipo,
    insertar_jugador,
    obtener_jugador_por_id,
    insertar_prueba,
    obtener_pruebas_por_jugador,
    obtener_resultados,
    obtener_todos_equipos,
    obtener_datos_grafica,
    obtener_todos_jugadores,
    eliminar_jugador,
    eliminar_temporada,
    eliminar_equipo,
    actualizar_jugador,
    obtener_categorias,
    obtener_grupos,
    obtener_jugadores_filtrados,
    insertar_tipo_prueba,
    obtener_pruebas_por_temporada,
    eliminar_tipo_prueba,
    obtener_temporada_de_jugador,
    insertar_resultado_prueba,
    obtener_resultados_dinamicos_jugador,
    obtener_resultados_dinamicos,
    obtener_datos_grafica_dinamicos,
    obtener_equipo_por_codigo,
    insertar_jugador_si_no_existe,
    insertar_equipo_si_no_existe
)

import os
from flask import Flask
import openpyxl
from werkzeug.utils import secure_filename

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

app = Flask(
    __name__,
    template_folder=os.path.join(BASE_DIR, "templates"),
    static_folder=os.path.join(BASE_DIR, "static")
)

UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

from webapp.db import crear_tablas
crear_tablas()

os.makedirs(os.path.join(BASE_DIR, "static", "graficas"), exist_ok=True)

def normalizar_nombre_equipo(nombre):
    nombre = str(nombre).strip().upper()

    conversiones = {
        "S.MASCULINO": "Senior Masc",
        "S.FEMENINO": "Senior Fem",
        "JUNIOR MASCULINO": "Junior Masc",
        "JUNIOR FEMENINO": "Junior Fem",
        "JUVENIL A": "Juvenil A",
        "JUVENIL B": "Juvenil B",
        "JUVENIL C": "Juvenil C",
        "CADETE A": "Cadete A",
        "CADETE B": "Cadete B",
        "CADETE C": "Cadete C",
        "INFANTIL A": "Infantil A",
        "INFANTIL B": "Infantil B",
        "INFANTIL C": "Infantil C",
        "ALEVIN A": "Alevín A",
        "ALEVIN B": "Alevín B",
        "ALEVIN C": "Alevín C",
        "BENJAMIN A": "Benjamín A",
        "BENJAMIN B": "Benjamín B",
        "PREBENJAMIN": "Prebenjamín"
    }

    return conversiones.get(nombre, nombre.title())

EQUIPOS_BASE = [
    {"codigo": 1, "nombre": "Senior Masc", "categoria": "Senior", "grupo": "Masc"},
    {"codigo": 2, "nombre": "Senior Fem", "categoria": "Senior", "grupo": "Fem"},
    {"codigo": 3, "nombre": "Junior Masc", "categoria": "Junior", "grupo": "Masc"},
    {"codigo": 4, "nombre": "Junior Fem", "categoria": "Junior", "grupo": "Fem"},
    {"codigo": 5, "nombre": "Juvenil A", "categoria": "Juvenil", "grupo": "A"},
    {"codigo": 6, "nombre": "Juvenil B", "categoria": "Juvenil", "grupo": "B"},
    {"codigo": 7, "nombre": "Juvenil C", "categoria": "Juvenil", "grupo": "C"},
    {"codigo": 8, "nombre": "Cadete A", "categoria": "Cadete", "grupo": "A"},
    {"codigo": 9, "nombre": "Cadete B", "categoria": "Cadete", "grupo": "B"},
    {"codigo": 10, "nombre": "Cadete C", "categoria": "Cadete", "grupo": "C"},
    {"codigo": 11, "nombre": "Infantil A", "categoria": "Infantil", "grupo": "A"},
    {"codigo": 12, "nombre": "Infantil B", "categoria": "Infantil", "grupo": "B"},
    {"codigo": 13, "nombre": "Infantil C", "categoria": "Infantil", "grupo": "C"},
    {"codigo": 14, "nombre": "Alevín A", "categoria": "Alevín", "grupo": "A"},
    {"codigo": 15, "nombre": "Alevín B", "categoria": "Alevín", "grupo": "B"},
    {"codigo": 16, "nombre": "Alevín C", "categoria": "Alevín", "grupo": "C"},
    {"codigo": 17, "nombre": "Benjamín A", "categoria": "Benjamín", "grupo": "A"},
    {"codigo": 18, "nombre": "Benjamín B", "categoria": "Benjamín", "grupo": "B"},
    {"codigo": 19, "nombre": "Prebenjamín", "categoria": "Prebenjamín", "grupo": ""},
]

def importar_excel_base(temporada_id):
    ruta_archivo = os.path.join(BASE_DIR, "data", "HOJA DE TEST PA APP CON PORTADA .xlsx")

    if not os.path.exists(ruta_archivo):
        print("⚠️ No se encontró el Excel base")
        return

    wb = openpyxl.load_workbook(ruta_archivo)

    # ==========================
    # 1) IMPORTAR EQUIPOS
    # ==========================
    hoja_equipos = wb["EQUIPOS"]

    for fila in hoja_equipos.iter_rows(min_row=2, values_only=True):
        codigo_equipo, nombre_equipo = fila[:2]

        if not codigo_equipo or not nombre_equipo:
            continue

        nombre_limpio = normalizar_nombre_equipo(nombre_equipo)

        insertar_equipo_si_no_existe(
            nombre_limpio,
            temporada_id,
            int(codigo_equipo)
        )

    # ==========================
    # 2) IMPORTAR JUGADORES
    # ==========================
    hoja_jugadores = wb["JUGADORES"]

    for fila in hoja_jugadores.iter_rows(min_row=2, values_only=True):
        _, nombre_jugador, codigo_equipo = fila[:3]

        if not nombre_jugador or not codigo_equipo:
            continue

        equipo = obtener_equipo_por_codigo(int(codigo_equipo), temporada_id)

        if equipo:
            insertar_jugador_si_no_existe(str(nombre_jugador).strip(), equipo[0])

def generar_graficas(jugador_id):
    datos = obtener_datos_grafica(jugador_id)

    if not datos:
        return None

    fechas = [d[0] for d in datos]
    sprint = [d[1] for d in datos]
    illinois = [d[2] for d in datos]
    yoyo = [d[3] for d in datos]

    carpeta_graficas = os.path.join(app.root_path, "static", "graficas")
    os.makedirs(carpeta_graficas, exist_ok=True)

    rutas = {}

    # --- Sprint 30m ---
    plt.figure(figsize=(10, 4))
    plt.plot(fechas, sprint, marker='o')
    plt.xlabel("Fecha")
    plt.ylabel("Tiempo (s)")
    plt.title("Evolución Sprint 30m")
    plt.xticks(rotation=45)
    plt.tight_layout()

    nombre_sprint = f"jugador_{jugador_id}_sprint.png"
    ruta_sprint = os.path.join(carpeta_graficas, nombre_sprint)
    plt.savefig(ruta_sprint)
    plt.close()
    rutas["sprint"] = f"static/graficas/{nombre_sprint}"

    # --- Illinois ---
    plt.figure(figsize=(10, 4))
    plt.plot(fechas, illinois, marker='o')
    plt.xlabel("Fecha")
    plt.ylabel("Tiempo (s)")
    plt.title("Evolución Illinois")
    plt.xticks(rotation=45)
    plt.tight_layout()

    nombre_illinois = f"jugador_{jugador_id}_illinois.png"
    ruta_illinois = os.path.join(carpeta_graficas, nombre_illinois)
    plt.savefig(ruta_illinois)
    plt.close()
    rutas["illinois"] = f"static/graficas/{nombre_illinois}"

    # --- YoYo ---
    plt.figure(figsize=(10, 4))
    plt.plot(fechas, yoyo, marker='o')
    plt.xlabel("Fecha")
    plt.ylabel("Resultado")
    plt.title("Evolución YoYo")
    plt.xticks(rotation=45)
    plt.tight_layout()

    nombre_yoyo = f"jugador_{jugador_id}_yoyo.png"
    ruta_yoyo = os.path.join(carpeta_graficas, nombre_yoyo)
    plt.savefig(ruta_yoyo)
    plt.close()
    rutas["yoyo"] = f"static/graficas/{nombre_yoyo}"

    return rutas


@app.route("/")
def inicio():
    return render_template("index.html")


@app.route("/gestion", methods=["GET", "POST"])
def gestion():
    if request.method == "POST":
        nombre = request.form["nombre_temporada"].strip()

        if nombre:
            temporada_id = insertar_temporada(nombre)
            importar_excel_base(temporada_id)

            return redirect(url_for("ver_temporada", temporada_id=temporada_id))

    temporadas = obtener_temporadas()
    return render_template("gestion.html", temporadas=temporadas)


@app.route("/temporada/<int:temporada_id>", methods=["GET", "POST"])
def ver_temporada(temporada_id):
    temporada = obtener_temporada_por_id(temporada_id)

    if request.method == "POST":

        # Crear equipo
        if "crear_equipo" in request.form:
            codigo_seleccionado = request.form.get("equipo_predefinido")

            if codigo_seleccionado:
                codigo_seleccionado = int(codigo_seleccionado)

                equipo_seleccionado = next(
                    (eq for eq in EQUIPOS_BASE if eq["codigo"] == codigo_seleccionado),
                    None
                )

                if equipo_seleccionado:
                    insertar_equipo_si_no_existe(
                        equipo_seleccionado["nombre"],
                        temporada_id,
                        equipo_seleccionado["codigo"]
                    )

            return redirect(url_for("ver_temporada", temporada_id=temporada_id))

        # Crear prueba física
        elif "crear_prueba" in request.form:
            nombre_prueba = request.form["nombre_prueba"].strip()
            unidad = request.form["unidad"].strip()

            if nombre_prueba:
                insertar_tipo_prueba(nombre_prueba, unidad, temporada_id)

    equipos = obtener_equipos_por_temporada(temporada_id)
    pruebas = obtener_pruebas_por_temporada(temporada_id)

    return render_template(
        "temporada.html",
        temporada=temporada,
        equipos=equipos,
        pruebas=pruebas,
        equipos_base=EQUIPOS_BASE
    )


@app.route("/equipo/<int:equipo_id>", methods=["GET", "POST"])
def ver_equipo(equipo_id):

    equipo = obtener_equipo_por_id(equipo_id)

    if request.method == "POST":
        nombre_jugador = request.form["nombre_jugador"].strip()

        if nombre_jugador:
            insertar_jugador(nombre_jugador, equipo_id)

        return redirect(url_for("ver_equipo", equipo_id=equipo_id))

    jugadores = obtener_jugadores_por_equipo(equipo_id)

    return render_template(
        "equipo.html",
        equipo=equipo,
        jugadores=jugadores
    )


@app.route("/jugador/<int:jugador_id>", methods=["GET", "POST"])
def ver_jugador(jugador_id):

    jugador = obtener_jugador_por_id(jugador_id)
    temporada = obtener_temporada_de_jugador(jugador_id)
    pruebas = obtener_pruebas_por_temporada(temporada[0])
    resultados = obtener_resultados_dinamicos_jugador(jugador_id)

    if request.method == "POST":
        fecha = request.form["fecha"]

        for prueba in pruebas:
            campo = f"prueba_{prueba[0]}"
            valor = request.form.get(campo)

            if valor and valor.strip() != "":
                insertar_resultado_prueba(jugador_id, prueba[0], fecha, float(valor))

        return redirect(url_for("ver_jugador", jugador_id=jugador_id))

    graficas = generar_graficas_jugador_dinamicas(jugador_id)

    return render_template(
        "jugador.html",
        jugador=jugador,
        pruebas=pruebas,
        resultados=resultados,
        graficas=graficas
    )

@app.route("/resultados")
def resultados():

    equipo_id = request.args.get("equipo_id")
    categoria = request.args.get("categoria")
    grupo = request.args.get("grupo")

    resultados = obtener_resultados_dinamicos(equipo_id, categoria, grupo)
    equipos = obtener_todos_equipos()
    categorias = obtener_categorias()
    grupos = obtener_grupos()

    return render_template(
        "resultados.html",
        resultados=resultados,
        equipos=equipos,
        categorias=categorias,
        grupos=grupos,
        categoria_actual=categoria,
        grupo_actual=grupo,
        equipo_actual=equipo_id
    )

def generar_graficas_comparacion(j1, j2):

    datos1 = obtener_datos_grafica(j1)
    datos2 = obtener_datos_grafica(j2)

    if not datos1 or not datos2:
        return None

    carpeta = os.path.join(app.root_path, "static", "graficas")
    os.makedirs(carpeta, exist_ok=True)

    rutas = {}

    def extraer(datos, index):
        return [d[0] for d in datos], [d[index] for d in datos]

    # --- SPRINT ---
    f1, s1 = extraer(datos1, 1)
    f2, s2 = extraer(datos2, 1)

    plt.figure(figsize=(10,4))
    plt.plot(f1, s1, marker='o', label="Jugador 1")
    plt.plot(f2, s2, marker='o', label="Jugador 2")
    plt.title("Comparación Sprint 30m")
    plt.legend()
    plt.xticks(rotation=45)
    plt.tight_layout()

    ruta = os.path.join(carpeta, "comp_sprint.png")
    plt.savefig(ruta)
    plt.close()
    rutas["sprint"] = "static/graficas/comp_sprint.png"

    # --- ILLINOIS ---
    f1, s1 = extraer(datos1, 2)
    f2, s2 = extraer(datos2, 2)

    plt.figure(figsize=(10,4))
    plt.plot(f1, s1, marker='o', label="Jugador 1")
    plt.plot(f2, s2, marker='o', label="Jugador 2")
    plt.title("Comparación Illinois")
    plt.legend()
    plt.xticks(rotation=45)
    plt.tight_layout()

    ruta = os.path.join(carpeta, "comp_illinois.png")
    plt.savefig(ruta)
    plt.close()
    rutas["illinois"] = "static/graficas/comp_illinois.png"

    # --- YOYO ---
    f1, s1 = extraer(datos1, 3)
    f2, s2 = extraer(datos2, 3)

    plt.figure(figsize=(10,4))
    plt.plot(f1, s1, marker='o', label="Jugador 1")
    plt.plot(f2, s2, marker='o', label="Jugador 2")
    plt.title("Comparación YoYo")
    plt.legend()
    plt.xticks(rotation=45)
    plt.tight_layout()

    ruta = os.path.join(carpeta, "comp_yoyo.png")
    plt.savefig(ruta)
    plt.close()
    rutas["yoyo"] = "static/graficas/comp_yoyo.png"

    return rutas

@app.route("/comparacion", methods=["GET", "POST"])
def comparacion():

    if request.method == "POST":
        categoria = request.form.get("categoria")
        grupo = request.form.get("grupo")
    else:
        categoria = request.args.get("categoria")
        grupo = request.args.get("grupo")

    jugadores = obtener_jugadores_filtrados(categoria, grupo)
    graficas = None

    if request.method == "POST":
        j1 = int(request.form["jugador1"])
        j2 = int(request.form["jugador2"])

        graficas = generar_graficas_comparacion_dinamicas(j1, j2)

    categorias = obtener_categorias()
    grupos = obtener_grupos()

    return render_template(
        "comparacion.html",
        jugadores=jugadores,
        graficas=graficas,
        categorias=categorias,
        grupos=grupos,
        categoria_actual=categoria,
        grupo_actual=grupo
    )

@app.route("/eliminar_jugador/<int:jugador_id>")
def borrar_jugador(jugador_id):

    jugador = obtener_jugador_por_id(jugador_id)

    if jugador:
        equipo_id = jugador[2]
        eliminar_jugador(jugador_id)
        return redirect(url_for("ver_equipo", equipo_id=equipo_id))

    return redirect(url_for("gestion"))

@app.route("/eliminar_equipo/<int:equipo_id>")
def borrar_equipo(equipo_id):

    equipo = obtener_equipo_por_id(equipo_id)

    if equipo:
        temporada_id = equipo[2]
        eliminar_equipo(equipo_id)
        return redirect(url_for("ver_temporada", temporada_id=temporada_id))

    return redirect(url_for("gestion"))

@app.route("/eliminar_temporada/<int:temporada_id>")
def borrar_temporada(temporada_id):

    eliminar_temporada(temporada_id)
    return redirect(url_for("gestion"))

@app.route("/editar_jugador/<int:jugador_id>", methods=["GET", "POST"])
def editar_jugador(jugador_id):

    jugador = obtener_jugador_por_id(jugador_id)

    if request.method == "POST":
        nuevo_nombre = request.form["nombre"].strip()

        if nuevo_nombre:
            actualizar_jugador(jugador_id, nuevo_nombre)

        return redirect(url_for("ver_equipo", equipo_id=jugador[2]))

    return render_template("editar_jugador.html", jugador=jugador)

@app.route("/eliminar_prueba/<int:prueba_id>/<int:temporada_id>")
def borrar_prueba(prueba_id, temporada_id):
    eliminar_tipo_prueba(prueba_id)
    return redirect(url_for("ver_temporada", temporada_id=temporada_id))

def generar_graficas_jugador_dinamicas(jugador_id):
    datos = obtener_datos_grafica_dinamicos(jugador_id)

    if not datos:
        return {}

    carpeta = os.path.join(app.root_path, "static", "graficas")
    os.makedirs(carpeta, exist_ok=True)

    pruebas = {}

    for nombre_prueba, unidad, fecha, valor in datos:
        if nombre_prueba not in pruebas:
            pruebas[nombre_prueba] = {"fechas": [], "valores": [], "unidad": unidad}

        pruebas[nombre_prueba]["fechas"].append(fecha)
        pruebas[nombre_prueba]["valores"].append(valor)

    rutas = {}

    for nombre_prueba, info in pruebas.items():
        plt.figure(figsize=(10, 4))
        plt.plot(info["fechas"], info["valores"], marker='o')
        plt.title(f"Evolución - {nombre_prueba}")
        plt.xlabel("Fecha")
        plt.ylabel(info["unidad"])
        plt.xticks(rotation=45)
        plt.tight_layout()

        nombre_archivo = f"jugador_{jugador_id}_{nombre_prueba.replace(' ', '_')}.png"
        ruta_completa = os.path.join(carpeta, nombre_archivo)
        plt.savefig(ruta_completa)
        plt.close()

        rutas[nombre_prueba] = f"static/graficas/{nombre_archivo}"

    return rutas

def generar_graficas_comparacion_dinamicas(j1, j2):

    datos1 = obtener_datos_grafica_dinamicos(j1)
    datos2 = obtener_datos_grafica_dinamicos(j2)

    carpeta = os.path.join(app.root_path, "static", "graficas")
    os.makedirs(carpeta, exist_ok=True)

    pruebas1 = {}
    pruebas2 = {}

    for nombre, unidad, fecha, valor in datos1:
        if nombre not in pruebas1:
            pruebas1[nombre] = {"fechas": [], "valores": [], "unidad": unidad}
        pruebas1[nombre]["fechas"].append(fecha)
        pruebas1[nombre]["valores"].append(valor)

    for nombre, unidad, fecha, valor in datos2:
        if nombre not in pruebas2:
            pruebas2[nombre] = {"fechas": [], "valores": [], "unidad": unidad}
        pruebas2[nombre]["fechas"].append(fecha)
        pruebas2[nombre]["valores"].append(valor)

    pruebas_comunes = set(pruebas1.keys()) & set(pruebas2.keys())
    rutas = {}

    for prueba in pruebas_comunes:
        plt.figure(figsize=(10, 4))
        plt.plot(pruebas1[prueba]["fechas"], pruebas1[prueba]["valores"], marker='o', label="Jugador 1")
        plt.plot(pruebas2[prueba]["fechas"], pruebas2[prueba]["valores"], marker='o', label="Jugador 2")
        plt.title(f"Comparación - {prueba}")
        plt.xlabel("Fecha")
        plt.ylabel(pruebas1[prueba]["unidad"])
        plt.legend()
        plt.xticks(rotation=45)
        plt.tight_layout()

        nombre_archivo = f"comparacion_{j1}_{j2}_{prueba.replace(' ', '_')}.png"
        ruta_completa = os.path.join(carpeta, nombre_archivo)
        plt.savefig(ruta_completa)
        plt.close()

        rutas[prueba] = f"static/graficas/{nombre_archivo}"

    return rutas

@app.route("/importar_jugadores/<int:temporada_id>", methods=["GET", "POST"])
def importar_jugadores(temporada_id):

    temporada = obtener_temporada_por_id(temporada_id)
    mensaje = None

    if request.method == "POST":
        archivo = request.files.get("archivo_excel")

        if archivo and archivo.filename.endswith(".xlsx"):
            nombre_seguro = secure_filename(archivo.filename)
            ruta_archivo = os.path.join(app.config["UPLOAD_FOLDER"], nombre_seguro)
            archivo.save(ruta_archivo)

            wb = openpyxl.load_workbook(ruta_archivo)

            # ==========================
            # 1) IMPORTAR EQUIPOS
            # ==========================
            hoja_equipos = wb["EQUIPOS"]

            equipos_importados = 0

            for fila in hoja_equipos.iter_rows(min_row=2, values_only=True):
                # Estructura esperada:
                # ID EQUIPO | EQUIPO
                codigo_equipo, nombre_equipo = fila[:2]

                if not codigo_equipo or not nombre_equipo:
                    continue

                nombre_limpio = normalizar_nombre_equipo(nombre_equipo)

                insertar_equipo_si_no_existe(
                    nombre_limpio,
                    temporada_id,
                    int(codigo_equipo)
                )
                equipos_importados += 1

            # ==========================
            # 2) IMPORTAR JUGADORES
            # ==========================
            hoja_jugadores = wb["JUGADORES"]

            jugadores_importados = 0
            no_encontrados = []

            for fila in hoja_jugadores.iter_rows(min_row=2, values_only=True):
                # Estructura esperada:
                # ID_Jugador | Nombre_Jugador | Equipo
                _, nombre_jugador, codigo_equipo = fila[:3]

                if not nombre_jugador or not codigo_equipo:
                    continue

                equipo = obtener_equipo_por_codigo(int(codigo_equipo), temporada_id)

                if equipo:
                    insertar_jugador_si_no_existe(str(nombre_jugador).strip(), equipo[0])
                    jugadores_importados += 1
                else:
                    no_encontrados.append((nombre_jugador, codigo_equipo))

            mensaje = (
                f"Importación completada: "
                f"{equipos_importados} equipos procesados y "
                f"{jugadores_importados} jugadores procesados."
            )

            return render_template(
                "importar_jugadores.html",
                temporada=temporada,
                mensaje=mensaje,
                no_encontrados=no_encontrados
            )

        else:
            mensaje = "Debes subir un archivo .xlsx válido"

    return render_template(
        "importar_jugadores.html",
        temporada=temporada,
        mensaje=mensaje,
        no_encontrados=[]
    )

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)